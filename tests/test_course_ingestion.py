import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document as DocxDocument

from ingestion.chunker import _word_windows
from ingestion.parsers import (
    DocumentMetadata,
    DocumentParseError,
    parse_document,
)
from ingestion.service import (
    CourseIngestionConfig,
    CourseIngestionError,
    ingest_course_package,
)


def write_package(
    root: Path,
    *,
    course_id: str = "political_philosophy",
    version: str = "1.0.0",
    document_name: str = "chapter-1.md",
    content: str | None = None,
    declare_source: bool = True,
) -> Path:
    documents = root / "documents"
    documents.mkdir(parents=True)
    document_path = documents / document_name
    document_path.write_text(
        content
        or (
            "# Chương 1\n"
            "Vật chất là một phạm trù triết học. "
            "Ý thức phản ánh thế giới khách quan.\n"
            "## Quan hệ\n"
            "Vật chất quyết định ý thức và ý thức tác động trở lại vật chất."
        ),
        encoding="utf-8",
    )
    source_block = ""
    if declare_source:
        source_block = "\n".join(
            [
                "sources:",
                f"  - path: documents/{document_name}",
                "    document_id: official_textbook",
                "    title: Giáo trình Triết học Mác - Lênin",
                "    source_type: official_textbook",
                "    source_authority: Nhà xuất bản kiểm thử",
                "    publication_year: 2025",
            ]
        )
    (root / "course.yaml").write_text(
        "\n".join(
            [
                f"course_id: {course_id}",
                "name: Triết học Mác - Lênin",
                "domain: political_theory",
                f"version: {version}",
                "language: vi",
                source_block,
            ]
        ),
        encoding="utf-8",
    )
    return root


def metadata_for(path: Path, source_path: str) -> DocumentMetadata:
    return DocumentMetadata(
        document_id="test_document",
        title="Tài liệu kiểm thử",
        source_path=source_path,
        source_type="reference",
        source_authority=None,
        publication_year=None,
        source_sha256="0" * 64,
        metadata_status="inferred",
    )


class DocumentParserTests(unittest.TestCase):
    def test_markdown_keeps_heading_and_line_citation(self):
        with tempfile.TemporaryDirectory(prefix="markdown-parser-") as directory:
            path = Path(directory) / "lesson.md"
            path.write_text(
                "# Chương 1\nNội dung một.\n## Mục 1\nNội dung hai.",
                encoding="utf-8",
            )
            sections = parse_document(
                path,
                metadata_for(path, "documents/lesson.md"),
            )

        self.assertEqual(len(sections), 2)
        self.assertEqual(sections[0].section_title, "Chương 1")
        self.assertEqual((sections[0].start_line, sections[0].end_line), (1, 2))
        self.assertIn("Nội dung hai", sections[1].text)

    def test_docx_keeps_heading_and_paragraph_citation(self):
        with tempfile.TemporaryDirectory(prefix="docx-parser-") as directory:
            path = Path(directory) / "lesson.docx"
            document = DocxDocument()
            document.add_heading("Chương 1", level=1)
            document.add_paragraph("Nội dung từ DOCX.")
            document.add_heading("Mục 1", level=2)
            document.add_paragraph("Nội dung tiếp theo.")
            document.save(path)

            sections = parse_document(
                path,
                metadata_for(path, "documents/lesson.docx"),
            )

        self.assertEqual(len(sections), 2)
        self.assertEqual(sections[0].section_title, "Chương 1")
        self.assertEqual(
            (sections[0].start_paragraph, sections[0].end_paragraph),
            (1, 2),
        )

    def test_pdf_keeps_real_page_number_and_rejects_empty_scan(self):
        class FakePage:
            def __init__(self, text):
                self.text = text

            def extract_text(self):
                return self.text

        class FakeReader:
            is_encrypted = False

            def __init__(self, pages):
                self.pages = pages

        with tempfile.TemporaryDirectory(prefix="pdf-parser-") as directory:
            path = Path(directory) / "lesson.pdf"
            path.write_bytes(b"%PDF-test")
            metadata = metadata_for(path, "documents/lesson.pdf")

            with patch(
                "ingestion.parsers.PdfReader",
                return_value=FakeReader([FakePage("Trang một"), FakePage("")]),
            ):
                sections = parse_document(path, metadata)
            self.assertEqual(len(sections), 1)
            self.assertEqual(sections[0].page_number, 1)

            with patch(
                "ingestion.parsers.PdfReader",
                return_value=FakeReader([FakePage("")]),
            ):
                with self.assertRaises(DocumentParseError):
                    parse_document(path, metadata)


class ChunkerTests(unittest.TestCase):
    def test_word_windows_do_not_create_overlap_only_tail(self):
        words = " ".join(f"w{index}" for index in range(10))
        windows = _word_windows(words, max_words=10, overlap_words=2)
        self.assertEqual(len(windows), 1)

        longer = " ".join(f"w{index}" for index in range(15))
        windows = _word_windows(longer, max_words=10, overlap_words=2)
        self.assertEqual(len(windows), 2)
        self.assertEqual(windows[1].split()[:2], ["w8", "w9"])


class CourseIngestionTests(unittest.TestCase):
    def test_ingests_package_with_grounded_citation_and_versioned_output(self):
        with tempfile.TemporaryDirectory(prefix="course-ingestion-") as directory:
            workspace = Path(directory)
            package = write_package(workspace / "package")
            output_root = workspace / "output"

            result = ingest_course_package(
                package,
                output_root=output_root,
                config=CourseIngestionConfig(max_words=12, overlap_words=3),
            )

            self.assertTrue(result.chunks_path.is_file())
            self.assertTrue(result.manifest_path.is_file())
            self.assertEqual(result.manifest.document_count, 1)
            self.assertGreaterEqual(result.manifest.section_count, 2)
            self.assertGreaterEqual(result.manifest.chunk_count, 2)
            self.assertEqual(
                result.output_directory,
                output_root.resolve()
                / "political_philosophy"
                / "1.0.0"
                / "processed",
            )

            first_chunk = result.chunks[0]
            self.assertEqual(first_chunk.course_id, "political_philosophy")
            self.assertEqual(first_chunk.course_version, "1.0.0")
            self.assertEqual(first_chunk.citation.document_id, "official_textbook")
            self.assertEqual(
                first_chunk.citation.source_authority,
                "Nhà xuất bản kiểm thử",
            )
            self.assertEqual(
                first_chunk.citation.source_path,
                "documents/chapter-1.md",
            )
            self.assertIsNotNone(first_chunk.citation.start_line)

            saved_chunks = json.loads(
                result.chunks_path.read_text(encoding="utf-8")
            )
            saved_manifest = json.loads(
                result.manifest_path.read_text(encoding="utf-8")
            )
            self.assertEqual(len(saved_chunks), result.manifest.chunk_count)
            self.assertEqual(
                saved_manifest["knowledge_version"],
                result.manifest.knowledge_version,
            )
            self.assertNotIn(
                str(package.resolve()),
                json.dumps(saved_chunks, ensure_ascii=False),
            )

    def test_output_is_immutable_without_force(self):
        with tempfile.TemporaryDirectory(prefix="course-ingestion-") as directory:
            workspace = Path(directory)
            package = write_package(workspace / "package")
            output_root = workspace / "output"
            first = ingest_course_package(package, output_root=output_root)

            with self.assertRaises(CourseIngestionError):
                ingest_course_package(package, output_root=output_root)

            overwritten = ingest_course_package(
                package,
                output_root=output_root,
                force=True,
            )
            self.assertEqual(
                first.manifest.knowledge_version,
                overwritten.manifest.knowledge_version,
            )

    def test_different_courses_have_isolated_ids_and_output(self):
        with tempfile.TemporaryDirectory(prefix="course-ingestion-") as directory:
            workspace = Path(directory)
            first_package = write_package(
                workspace / "first",
                course_id="political_philosophy",
                declare_source=False,
            )
            second_package = write_package(
                workspace / "second",
                course_id="hochiminh_thought",
                declare_source=False,
            )
            output_root = workspace / "output"

            first = ingest_course_package(
                first_package,
                output_root=output_root,
            )
            second = ingest_course_package(
                second_package,
                output_root=output_root,
            )

            self.assertNotEqual(first.chunks[0].chunk_id, second.chunks[0].chunk_id)
            self.assertNotEqual(
                first.manifest.knowledge_version,
                second.manifest.knowledge_version,
            )
            self.assertNotEqual(first.output_directory, second.output_directory)
            self.assertEqual(first.chunks[0].course_id, "political_philosophy")
            self.assertEqual(second.chunks[0].course_id, "hochiminh_thought")


if __name__ == "__main__":
    unittest.main()
