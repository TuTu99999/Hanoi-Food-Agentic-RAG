import hashlib
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader


MAX_DOCUMENT_BYTES = 50 * 1024 * 1024


class DocumentParseError(ValueError):
    pass


@dataclass(frozen=True)
class DocumentMetadata:
    document_id: str
    title: str
    source_path: str
    source_type: str
    source_authority: str | None
    publication_year: int | None
    source_sha256: str
    metadata_status: str


@dataclass(frozen=True)
class ParsedSection:
    document: DocumentMetadata
    section_id: str
    section_title: str | None
    text: str
    page_number: int | None = None
    start_line: int | None = None
    end_line: int | None = None
    start_paragraph: int | None = None
    end_paragraph: int | None = None


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def slugify(value: str, fallback: str = "document") -> str:
    ascii_value = unicodedata.normalize("NFKD", value).encode(
        "ascii",
        "ignore",
    ).decode("ascii")
    slug = re.sub(r"[^a-zA-Z0-9]+", "_", ascii_value).strip("_").lower()
    return slug or fallback


def automatic_document_id(source_path: str) -> str:
    stem = slugify(Path(source_path).stem)
    suffix = hashlib.sha256(source_path.encode("utf-8")).hexdigest()[:8]
    return f"{stem}_{suffix}"[:80]


def _clean_block(lines: list[str]) -> str:
    cleaned_lines = [
        re.sub(r"[ \t]+", " ", line).strip()
        for line in lines
    ]
    return "\n".join(line for line in cleaned_lines if line).strip()


def _section_id(document_id: str, locator: str) -> str:
    suffix = hashlib.sha256(locator.encode("utf-8")).hexdigest()[:10]
    return f"{document_id}_section_{suffix}"


def _heading_from_line(line: str) -> str | None:
    stripped = line.strip()
    markdown_match = re.match(r"^#{1,6}\s+(.+?)\s*#*$", stripped)
    if markdown_match:
        return markdown_match.group(1).strip()
    if re.match(r"^(chương|phần|bài)\s+[0-9ivxlcdm]+\b", stripped, re.I):
        return stripped
    return None


def _parse_text_document(
    path: Path,
    metadata: DocumentMetadata,
) -> list[ParsedSection]:
    try:
        text = path.read_text(encoding="utf-8-sig")
    except (OSError, UnicodeDecodeError) as exc:
        raise DocumentParseError(
            f"Không thể đọc tài liệu UTF-8: {metadata.source_path}"
        ) from exc

    lines = text.splitlines()
    sections: list[ParsedSection] = []
    current_title: str | None = None
    current_lines: list[str] = []
    start_line = 1

    def flush(end_line: int) -> None:
        nonlocal current_lines
        content = _clean_block(current_lines)
        if not content:
            current_lines = []
            return
        locator = f"lines:{start_line}-{end_line}:{current_title or ''}"
        sections.append(
            ParsedSection(
                document=metadata,
                section_id=_section_id(metadata.document_id, locator),
                section_title=current_title,
                text=content,
                start_line=start_line,
                end_line=max(start_line, end_line),
            )
        )
        current_lines = []

    for line_number, line in enumerate(lines, start=1):
        heading = _heading_from_line(line)
        if heading is not None:
            flush(line_number - 1)
            current_title = heading
            current_lines = [heading]
            start_line = line_number
        else:
            if not current_lines:
                start_line = line_number
            current_lines.append(line)
    flush(len(lines))
    return sections


def _parse_pdf_document(
    path: Path,
    metadata: DocumentMetadata,
) -> list[ParsedSection]:
    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted and reader.decrypt("") == 0:
            raise DocumentParseError(
                f"PDF được bảo vệ bằng mật khẩu: {metadata.source_path}"
            )
        sections = []
        for page_number, page in enumerate(reader.pages, start=1):
            text = _clean_block((page.extract_text() or "").splitlines())
            if not text:
                continue
            locator = f"page:{page_number}"
            sections.append(
                ParsedSection(
                    document=metadata,
                    section_id=_section_id(metadata.document_id, locator),
                    section_title=f"Trang {page_number}",
                    text=text,
                    page_number=page_number,
                )
            )
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError(
            f"Không thể đọc PDF: {metadata.source_path}"
        ) from exc
    return sections


def _parse_docx_document(
    path: Path,
    metadata: DocumentMetadata,
) -> list[ParsedSection]:
    try:
        document = DocxDocument(str(path))
    except Exception as exc:
        raise DocumentParseError(
            f"Không thể đọc DOCX: {metadata.source_path}"
        ) from exc

    sections: list[ParsedSection] = []
    current_title: str | None = None
    current_lines: list[str] = []
    start_paragraph = 1

    def flush(end_paragraph: int) -> None:
        nonlocal current_lines
        content = _clean_block(current_lines)
        if not content:
            current_lines = []
            return
        locator = (
            f"paragraphs:{start_paragraph}-{end_paragraph}:"
            f"{current_title or ''}"
        )
        sections.append(
            ParsedSection(
                document=metadata,
                section_id=_section_id(metadata.document_id, locator),
                section_title=current_title,
                text=content,
                start_paragraph=start_paragraph,
                end_paragraph=max(start_paragraph, end_paragraph),
            )
        )
        current_lines = []

    for paragraph_number, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        style_name = (paragraph.style.name or "").casefold()
        is_heading = style_name.startswith("heading") and bool(text)
        if is_heading:
            flush(paragraph_number - 1)
            current_title = text
            current_lines = [text]
            start_paragraph = paragraph_number
        else:
            if not current_lines:
                start_paragraph = paragraph_number
            current_lines.append(text)
    flush(len(document.paragraphs))
    return sections


def parse_document(
    path: Path,
    metadata: DocumentMetadata,
) -> list[ParsedSection]:
    try:
        file_size = path.stat().st_size
    except OSError as exc:
        raise DocumentParseError(
            f"Không thể đọc file: {metadata.source_path}"
        ) from exc
    if file_size > MAX_DOCUMENT_BYTES:
        raise DocumentParseError(
            f"Tài liệu vượt quá 50 MB: {metadata.source_path}"
        )

    suffix = path.suffix.casefold()
    if suffix in {".txt", ".md"}:
        sections = _parse_text_document(path, metadata)
    elif suffix == ".pdf":
        sections = _parse_pdf_document(path, metadata)
    elif suffix == ".docx":
        sections = _parse_docx_document(path, metadata)
    else:
        raise DocumentParseError(
            f"Định dạng tài liệu chưa được hỗ trợ: {metadata.source_path}"
        )

    if not sections:
        raise DocumentParseError(
            "Không trích xuất được nội dung từ tài liệu. "
            f"PDF scan cần OCR ở giai đoạn sau: {metadata.source_path}"
        )
    return sections
